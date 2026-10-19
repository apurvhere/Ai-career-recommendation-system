from io import BytesIO
from docx import Document
import fitz
from rest_framework.decorators import api_view
from rest_framework.response import Response
from .engine import detect_skills,recommend,skill_gap,load_careers
from .serializers import ResumeUploadSerializer,SkillsSerializer,SkillGapSerializer,AnalysisRecordSerializer
from .models import AnalysisRecord
@api_view(["GET"])
def health(request): return Response({"status":"ok","message":"Career Recommendation API is running."})
@api_view(["GET"])
def careers(request): return Response(load_careers())
@api_view(["POST"])
def parse_resume(request):
    ser=ResumeUploadSerializer(data=request.data)
    if not ser.is_valid(): return Response(ser.errors,status=400)
    f=ser.validated_data["resume"]; data=f.read()
    try:
        if f.name.lower().endswith(".pdf"):
            with fitz.open(stream=data,filetype="pdf") as doc: text="\n".join(p.get_text() for p in doc)
        else:
            doc=Document(BytesIO(data)); text="\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    except Exception as exc: return Response({"error":f"Could not read document: {exc}"},status=400)
    if not text.strip(): return Response({"error":"No selectable text found. This may be a scanned PDF; OCR is not included."},status=400)
    return Response({"filename":f.name,"text":text[:20000],"skills":detect_skills(text),"message":"Review detected skills before analysis."})
@api_view(["POST"])
def recommendations(request):
    ser=SkillsSerializer(data=request.data)
    if not ser.is_valid(): return Response(ser.errors,status=400)
    return Response({"recommendations":recommend(ser.validated_data["skills"]),"skills_used":ser.validated_data["skills"],"note":"Compatibility score is skill overlap, not a job probability."})
@api_view(["POST"])
def gaps(request):
    ser=SkillGapSerializer(data=request.data)
    if not ser.is_valid(): return Response(ser.errors,status=400)
    result=skill_gap(ser.validated_data["skills"],ser.validated_data["career"])
    return Response(result if result else {"error":"Career not found."},status=200 if result else 404)
@api_view(["GET","POST"])
def history(request):
    if request.method=="GET": return Response(AnalysisRecordSerializer(AnalysisRecord.objects.order_by("-created_at")[:50],many=True).data)
    ser=AnalysisRecordSerializer(data=request.data)
    if ser.is_valid(): return Response(AnalysisRecordSerializer(ser.save()).data,status=201)
    return Response(ser.errors,status=400)
@api_view(["DELETE"])
def delete_history(request,record_id):
    try: AnalysisRecord.objects.get(pk=record_id).delete()
    except AnalysisRecord.DoesNotExist: return Response({"error":"Record not found."},status=404)
    return Response(status=204)
